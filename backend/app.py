import os, json, re, mimetypes, io, base64, hashlib, time
from datetime import datetime
from zoneinfo import ZoneInfo
from pathlib import Path

import httpx
from flask import Flask, request, jsonify, Response, stream_with_context
from flask_cors import CORS
from dotenv import load_dotenv
from openai import OpenAI

BASE_DIR=Path(__file__).resolve().parent
APP_VERSION='v1'
load_dotenv(BASE_DIR/'.env')
GROQ_API_KEY=os.getenv('GROQ_API_KEY','').strip()
GEMINI_API_KEY=os.getenv('GEMINI_API_KEY','').strip()
GROQ_BASE='https://api.groq.com/openai/v1'
GROQ_MODEL=os.getenv('GROQ_MODEL','openai/gpt-oss-20b')
GEMINI_MODEL=os.getenv('GEMINI_MODEL','gemini-3.8-flash')
GEMINI_VISION_FALLBACKS=[x.strip() for x in os.getenv('GEMINI_VISION_FALLBACKS','gemini-3.8-flash,gemini-3.7-flash,gemini-3.6-flash,gemini-3.5-flash').split(',') if x.strip()]
if GEMINI_MODEL not in GEMINI_VISION_FALLBACKS: GEMINI_VISION_FALLBACKS.insert(0,GEMINI_MODEL)
GEMINI_IMAGE_MODEL=os.getenv('GEMINI_IMAGE_MODEL','gemini-3.1-flash-image')
GEMINI_VIDEO_MODEL=os.getenv('GEMINI_VIDEO_MODEL','veo-3.1-generate-preview')
FRONTEND_ORIGIN=os.getenv('FRONTEND_ORIGIN','https://vaani-ai-workspace-frontend.vercel.app')

app=Flask(__name__)
app.config['MAX_CONTENT_LENGTH']=20*1024*1024
allowed=[x.strip() for x in FRONTEND_ORIGIN.split(',') if x.strip()]
allowed += ['http://127.0.0.1:5500','http://localhost:5500','http://127.0.0.1:5288','http://localhost:5288']
CORS(app, origins=allowed or '*')
groq_client=OpenAI(api_key=GROQ_API_KEY,base_url=GROQ_BASE,timeout=90.0) if GROQ_API_KEY else None
try:
    from google import genai
    from google.genai import types as genai_types
    gemini_client=genai.Client(api_key=GEMINI_API_KEY) if GEMINI_API_KEY else None
except Exception:
    genai=None;genai_types=None;gemini_client=None

MODE_PROMPTS={
 'General':'Answer naturally and helpfully. Handle general knowledge, practical questions, explanations and everyday tasks.',
 'Coding Assistant':'Act as a careful software engineer. Explain architecture, code, debugging, APIs, databases and testing clearly. Prefer runnable examples.',
 'Interview Prep':'Act as an interview coach. Give accurate answers, examples, follow-up questions and practical interview phrasing.',
 'Explain Simply':'Explain difficult topics in simple language with analogies and step-by-step structure. Avoid unnecessary jargon.'}
BASE_PROMPT='''You are Vaani, an AI workspace assistant.\nBe useful, direct, honest and readable. Use the user language naturally, including Hinglish when appropriate.\nNever invent current facts such as weather, time, prices or scores. If live context is supplied, use it and identify it as live context.\nFor uploaded files, use only the supplied extracted content and visual inputs. If the file does not contain enough information, say so.\nDo not claim to have generated, edited, executed, searched or verified something unless the system actually supplied the result.\nPrefer structured answers with short paragraphs, bullets and code blocks when useful.'''
MODEL_FALLBACK=GROQ_MODEL
SUPPORTED_MODELS={GROQ_MODEL}
IMAGE_EXTENSIONS={'.png','.jpg','.jpeg','.webp'}
TEXT_EXTENSIONS={'.txt','.md','.csv','.tsv','.json','.py','.js','.ts','.jsx','.tsx','.html','.css','.java','.cpp','.c','.h','.hpp','.cs','.go','.rs','.php','.sql','.xml','.yaml','.yml','.sh','.bat','.ps1'}
DATA_EXTENSIONS={'.csv','.tsv','.xlsx','.xlsm','.xls'}
MAX_TEXT=180000
MAX_INLINE_FILE_BYTES=18*1024*1024


def now_ist(): return datetime.now(ZoneInfo('Asia/Kolkata')).isoformat()
def sha256(data): return hashlib.sha256(data).hexdigest()[:16]
def json_error(message,status=400,code='bad_request'): return jsonify({'ok':False,'error':message,'code':code}),status

def weather_code_summary(code):
    return {0:'Clear sky',1:'Mainly clear',2:'Partly cloudy',3:'Overcast',45:'Fog',48:'Rime fog',51:'Light drizzle',53:'Drizzle',55:'Heavy drizzle',61:'Light rain',63:'Rain',65:'Heavy rain',71:'Light snow',73:'Snow',75:'Heavy snow',80:'Rain showers',81:'Rain showers',82:'Heavy showers',95:'Thunderstorm',96:'Thunderstorm with hail',99:'Thunderstorm with hail'}.get(code,'Current conditions')

async def get_weather_coords(lat,lon,label='Your location',country='',http=None):
    own=http is None; h=http or httpx.AsyncClient(timeout=12)
    try:
        r=await h.get('https://api.open-meteo.com/v1/forecast',params={'latitude':lat,'longitude':lon,'current':'temperature_2m,relative_humidity_2m,apparent_temperature,weather_code,wind_speed_10m','timezone':'auto'});r.raise_for_status();d=r.json();c=d['current']
        return {'ok':True,'city':label,'country':country,'timezone':d.get('timezone'),'latitude':lat,'longitude':lon,'temperature':c['temperature_2m'],'feels_like':c['apparent_temperature'],'humidity':c['relative_humidity_2m'],'wind':c['wind_speed_10m'],'summary':weather_code_summary(c['weather_code']),'time':c['time']}
    finally:
        if own: await h.aclose()

async def get_weather(city):
    async with httpx.AsyncClient(timeout=12) as h:
        g=await h.get('https://geocoding-api.open-meteo.com/v1/search',params={'name':city,'count':1,'language':'en','format':'json'});g.raise_for_status();rs=g.json().get('results') or []
        if not rs: raise ValueError('Location not found')
        loc=rs[0];return await get_weather_coords(loc['latitude'],loc['longitude'],loc.get('name',city),loc.get('country',''),h)

def build_context(live):
    c=[f"Current server time in India (IST): {now_ist()}"]
    if isinstance(live,dict):
        if live.get('timezone'): c.append('Client timezone: '+str(live['timezone']))
        if live.get('weather'): c.append('Live weather: '+json.dumps(live['weather'],ensure_ascii=False))
        if live.get('workspace'): c.append('Active workspace: '+str(live['workspace']))
        if live.get('memory'): c.append('Explicit user memory notes: '+json.dumps(live['memory'],ensure_ascii=False)[:7000])
        if live.get('files'): c.append('Uploaded file text: '+str(live['files'])[:30000])
    return '\n'.join(c)

def build_text_messages(message,history,mode,live):
    sys=BASE_PROMPT+'\nMode: '+MODE_PROMPTS.get(mode,MODE_PROMPTS['General'])+'\n\nLIVE CONTEXT:\n'+build_context(live)
    msgs=[{'role':'system','content':sys}]
    for x in (history or [])[-12:]:
        if x.get('role') in ('user','assistant') and x.get('content'): msgs.append({'role':x['role'],'content':str(x['content'])[:12000]})
    msgs.append({'role':'user','content':str(message)[:20000]})
    return msgs

def has_attachments(live):
    return bool((live or {}).get('images') or (live or {}).get('attachments'))

def data_url_bytes(data_url):
    if not isinstance(data_url,str) or not data_url.startswith('data:') or ';base64,' not in data_url: raise ValueError('Invalid attachment data')
    header,b64=data_url.split(';base64,',1);mime=header[5:] or 'application/octet-stream'
    return base64.b64decode(b64),mime

def _visual_items(live):
    attachments=(live or {}).get('attachments') or []
    images=(live or {}).get('images') or []
    items=[]; seen=set()
    for item in images+attachments:
        if not isinstance(item,dict):
            continue
        data_url=item.get('data_url')
        if not data_url or data_url in seen:
            continue
        seen.add(data_url)
        data,mime=data_url_bytes(data_url)
        if len(data)>MAX_INLINE_FILE_BYTES:
            continue
        items.append({'name':item.get('name') or 'attachment','data':data,'mime':mime,'b64':data_url.split(';base64,',1)[1]})
    return items

def visual_prompt(message,mode):
    return (BASE_PROMPT + '\nMode: ' + MODE_PROMPTS.get(mode,MODE_PROMPTS['General']) +
            '\n\nIMPORTANT VISUAL-INSTRUCTION:\n'
            'The user has attached one or more files and the attachment bytes are present in this request. '
            'You MUST inspect the actual visual content before answering. Do not ask the user to describe the image, '
            'do not say that you cannot see images, and do not give a generic request for a rundown. '
            'If it is a screenshot, read and explain the visible UI, text, code, error, diagram, or other content as appropriate. '
            'Use only details supported by the attachment. If some text is unreadable, say exactly what part is unclear. '
            'Answer the user directly based on what you can observe.\n\nUSER REQUEST:\n' + str(message)[:20000])

def gemini_contents(message,history,mode,live):
    if not gemini_client or not GEMINI_API_KEY:
        raise RuntimeError('GEMINI_API_KEY is not configured')
    visual_items=_visual_items(live)
    if not visual_items:
        if has_attachments(live):
            raise RuntimeError('The image could not be attached to the vision request. Please reattach it and try again.')
        context=build_context(live)
        prior=[]
        for x in (history or [])[-8:]:
            if x.get('role') in ('user','assistant') and x.get('content'):
                prior.append(f"{x['role'].upper()}: {str(x['content'])[:5000]}")
        prompt=(BASE_PROMPT+'\nMode: '+MODE_PROMPTS.get(mode,MODE_PROMPTS['General'])+'\n'+context+
                '\n\nRECENT CHAT:\n'+('\\n'.join(prior) if prior else '(none)')+
                '\n\nUSER REQUEST:\n'+str(message)[:20000])
        return [genai_types.Content(role='user',parts=[genai_types.Part.from_text(text=prompt)])]
    user_parts=[genai_types.Part.from_bytes(data=x['data'],mime_type=x['mime']) for x in visual_items]
    user_parts.append(genai_types.Part.from_text(text=visual_prompt(message,mode)))
    return [genai_types.Content(role='user',parts=user_parts)]

def gemini_interactions_input(message,mode,live):
    visual_items=_visual_items(live)
    if not visual_items:
        raise RuntimeError('No visual attachment was supplied.')
    inp=[{'type':'text','text':visual_prompt(message,mode)}]
    for x in visual_items:
        if not x['mime'].startswith('image/'):
            raise RuntimeError('This attachment type uses the standard Gemini vision route.')
        inp.append({'type':'image','data':x['b64'],'mime_type':x['mime']})
    return inp, len(visual_items)

def groq_complete(messages,stream=False):
    if not GROQ_API_KEY or groq_client is None: raise RuntimeError('GROQ_API_KEY is not configured. Add it to backend/.env and restart the server.')
    return groq_client.chat.completions.create(model=GROQ_MODEL,messages=messages,temperature=0.45,max_completion_tokens=1800,reasoning_effort='low',stream=stream)

def _gemini_error_code(exc):
    return getattr(exc,'code',None) or getattr(exc,'status_code',None) or getattr(getattr(exc,'response',None),'status_code',None)

def gemini_complete(message,history,mode,live):
    if not gemini_client:
        raise RuntimeError('Vision is temporarily unavailable.')
    retryable={429,500,502,503,504}
    failures=[]
    has_visual=has_attachments(live)

    # For image requests use the current Gemini Interactions multimodal format first.
    # It sends the actual base64 image bytes as an image input, independently of chat history.
    if has_visual and hasattr(gemini_client,'interactions'):
        try:
            inp,count=gemini_interactions_input(message or 'Inspect the attached image and explain what is visible.',mode,live)
            for model in GEMINI_VISION_FALLBACKS:
                try:
                    interaction=gemini_client.interactions.create(model=model,input=inp)
                    text=getattr(interaction,'output_text',None) or ''
                    if text.strip():
                        return text.strip(),model
                    failures.append(f'{model}:empty_interaction_response')
                except Exception as e:
                    code=_gemini_error_code(e);failures.append(f'{model}:interaction:{code or type(e).__name__}')
                    app.logger.warning('Vision interaction model %s failed (%s); trying next model',model,code or type(e).__name__)
                    if code not in retryable:
                        break
                    time.sleep(0.35)
        except Exception as e:
            failures.append(f'interactions:{type(e).__name__}')
            app.logger.warning('Gemini interaction vision path unavailable: %s; using generateContent fallback',type(e).__name__)

    config=genai_types.GenerateContentConfig(max_output_tokens=2400,thinking_config=genai_types.ThinkingConfig(thinking_level='low'))
    contents=gemini_contents(message,history,mode,live)
    for model in GEMINI_VISION_FALLBACKS:
        try:
            response=gemini_client.models.generate_content(model=model,contents=contents,config=config)
            text=getattr(response,'text',None) or ''
            if text.strip(): return text.strip(),model
            failures.append(f'{model}:empty_response')
        except Exception as e:
            code=_gemini_error_code(e);failures.append(f'{model}:{code or type(e).__name__}')
            app.logger.warning('Vision model %s failed (%s); trying fallback if available',model,code or type(e).__name__)
            if code not in retryable: break
            time.sleep(0.35)
    app.logger.error('All vision paths failed: %s','; '.join(failures))
    raise RuntimeError('Image analysis is temporarily unavailable. Please try again in a moment.')

def extract_city(text):
    for p in [r'(?:weather|temperature|forecast)\s+(?:in|at|of)\s+([A-Za-z .\'-]+)',r'([A-Za-z .\'-]+)\s+(?:ka|ki|ke)\s+(?:weather|temperature|mausam)',r'(?:weather|temperature|mausam)\s+(?:of|in)?\s*([A-Za-z .\'-]+)']:
        m=re.search(p,text,re.I)
        if m:return m.group(1).strip(' .?!,')
    return None

@app.get('/')
def root(): return jsonify({'app':'Vaani','version':APP_VERSION,'status':'ok','groq_configured':bool(GROQ_API_KEY),'gemini_configured':bool(GEMINI_API_KEY)})
@app.get('/api/health')
def health(): return jsonify({'ok':True,'version':APP_VERSION,'groq_configured':bool(GROQ_API_KEY),'gemini_configured':bool(GEMINI_API_KEY),'vision_model':GEMINI_MODEL,'vision_models':GEMINI_VISION_FALLBACKS,'text_model':GROQ_MODEL,'time_ist':now_ist()})
@app.get('/api/time')
def api_time():
    city=request.args.get('city','India');zones={'india':'Asia/Kolkata','delhi':'Asia/Kolkata','mumbai':'Asia/Kolkata','london':'Europe/London','new york':'America/New_York','tokyo':'Asia/Tokyo','dubai':'Asia/Dubai','singapore':'Asia/Singapore','sydney':'Australia/Sydney'};zone=zones.get(city.lower(),'Asia/Kolkata');d=datetime.now(ZoneInfo(zone));return jsonify({'ok':True,'city':city,'timezone':zone,'time':d.isoformat(),'formatted':d.strftime('%d %b %Y, %I:%M:%S %p')})
@app.get('/api/weather')
def api_weather():
    import asyncio
    try:
        lat=request.args.get('lat');lon=request.args.get('lon')
        if lat is not None and lon is not None:
            a,b=float(lat),float(lon)
            if not(-90<=a<=90 and -180<=b<=180): raise ValueError('Invalid coordinates')
            return jsonify(asyncio.run(get_weather_coords(a,b)))
        city=request.args.get('city','New Delhi').strip() or 'New Delhi';return jsonify(asyncio.run(get_weather(city)))
    except Exception as e:return json_error(str(e),502,'weather_error')

def search_web(query, limit=10):
    """Fetch general web results with resilient DuckDuckGo HTML/lite fallbacks.
    The parser deliberately supports both current result layouts so a UI empty state
    is not caused merely by a selector/layout change on the provider side.
    """
    from bs4 import BeautifulSoup
    headers={'User-Agent':'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Vaani/6.4.1 Chrome/154 Safari/537.36','Accept-Language':'en-IN,en;q=0.9'}
    endpoints=[
        ('DuckDuckGo HTML','https://html.duckduckgo.com/html/'),
        ('DuckDuckGo Lite','https://lite.duckduckgo.com/lite/'),
    ]
    errors=[]
    with httpx.Client(timeout=15,headers=headers,follow_redirects=True) as h:
        for source,url in endpoints:
            try:
                r=h.get(url,params={'q':query,'kl':'in-en'}); r.raise_for_status()
                soup=BeautifulSoup(r.text,'html.parser'); items=[]
                # Standard HTML layout
                for res in soup.select('.result, .results_links_deep'):
                    a=res.select_one('.result__a, a.result-link, .result-link')
                    sn=res.select_one('.result__snippet, .result-snippet')
                    if not a: continue
                    href=(a.get('href') or '').strip(); title=a.get_text(' ',strip=True); text=sn.get_text(' ',strip=True) if sn else ''
                    if href and title and href.startswith(('http://','https://')) and not any(x['url']==href for x in items):
                        items.append({'title':title,'text':text,'url':href})
                # Lite layout fallback: result links/snippets are not wrapped in .result
                if not items:
                    links=soup.select('a.result-link, a.result__a')
                    for a in links:
                        href=(a.get('href') or '').strip(); title=a.get_text(' ',strip=True)
                        parent=a.parent.parent if a.parent and a.parent.parent else a.parent
                        sn=parent.select_one('.result-snippet, .result__snippet') if parent else None
                        text=sn.get_text(' ',strip=True) if sn else ''
                        if href and title and href.startswith(('http://','https://')) and not any(x['url']==href for x in items):
                            items.append({'title':title,'text':text,'url':href})
                if items: return items[:limit], source
                errors.append(f'{source}: no parseable results')
            except Exception as e:
                errors.append(f'{source}: {type(e).__name__}')
    raise RuntimeError('Web search provider returned no usable results. ' + '; '.join(errors))

@app.get('/api/search')
def api_search():
    q=request.args.get('q','').strip()
    if not q:return json_error('Missing query')
    try:
        items,source=search_web(q,10)
        return jsonify({'ok':True,'query':q,'results':items,'source':source})
    except Exception as e:return json_error(str(e),502,'search_error')

def validate_image_bytes(data,mime,ext):
    if ext not in IMAGE_EXTENSIONS and not mime.startswith('image/'): return False
    if data.startswith(b'\x89PNG\r\n\x1a\n'): return mime in ('image/png','application/octet-stream')
    if data.startswith(b'\xff\xd8\xff'): return mime in ('image/jpeg','image/jpg','application/octet-stream')
    if data[:12]==b'RIFF'+data[4:8]+b'WEBP': return mime in ('image/webp','application/octet-stream')
    return False

def extract_docx(data):
    from docx import Document
    doc=Document(io.BytesIO(data));parts=[p.text for p in doc.paragraphs if p.text.strip()]
    for ti,table in enumerate(doc.tables,1):
        rows=[]
        for row in table.rows: rows.append(' | '.join(cell.text.replace('\n',' ') for cell in row.cells))
        if rows:parts.append(f'Table {ti}:\n'+'\n'.join(rows))
    return '\n\n'.join(parts)

@app.post('/api/files/analyze')
def analyze_file():
    f=request.files.get('file')
    if not f:return json_error('No file provided')
    data=f.read();name=Path(f.filename or 'file').name;ext=Path(name).suffix.lower();mime=f.mimetype or mimetypes.guess_type(name)[0] or 'application/octet-stream'
    if len(data)>20*1024*1024:return json_error('File exceeds the 20 MB upload limit',413,'file_too_large')
    try:
        if ext in IMAGE_EXTENSIONS or mime.startswith('image/'):
            if ext not in IMAGE_EXTENSIONS or not validate_image_bytes(data,mime,ext):return json_error('Supported images are valid PNG, JPG/JPEG or WebP files.',400,'invalid_image')
            b64=base64.b64encode(data).decode('ascii');return jsonify({'ok':True,'name':name,'type':ext[1:],'mime':mime if mime.startswith('image/') else ('image/jpeg' if ext in {'.jpg','.jpeg'} else 'image/'+ext[1:]),'kind':'image','text':'Image uploaded. Ask Vaani to describe, inspect or reason about what is visible.','data_url':f"data:{mime if mime.startswith('image/') else 'image/jpeg'};base64,{b64}",'characters':0,'size':len(data),'sha256':sha256(data)})
        if ext in TEXT_EXTENSIONS:text=data.decode('utf-8','ignore')
        elif ext=='.pdf':
            from pypdf import PdfReader
            reader=PdfReader(io.BytesIO(data));pages=[]
            for i,p in enumerate(reader.pages,1): pages.append(f'--- Page {i} ---\n{p.extract_text() or ""}')
            text='\n'.join(pages)
            if not text.strip(): text='[PDF contains no extractable text. It may be scanned/image-only; Gemini can inspect the PDF when the inline PDF payload is available.]'
        elif ext in {'.xlsx','.xlsm','.xls'}:
            import pandas as pd
            book=pd.ExcelFile(io.BytesIO(data));parts=[]
            for sheet in book.sheet_names:
                df=book.parse(sheet);parts.append(f'Sheet: {sheet}\nRows: {len(df)} | Columns: {len(df.columns)}\nColumns: {list(map(str,df.columns))}\nPreview:\n{df.head(40).to_csv(index=False)}')
            text='\n\n'.join(parts)
        elif ext=='.docx':text=extract_docx(data)
        else:return json_error(f'Unsupported file type: {ext or mime}',400,'unsupported_file')
        text=text[:MAX_TEXT];payload={'ok':True,'name':name,'type':ext[1:] if ext else mime,'mime':mime,'kind':'document','text':text,'characters':len(text),'size':len(data),'sha256':sha256(data),'note':'Gemini will use the native PDF when the file is small enough for inline transfer.' if ext=='.pdf' and len(data)<=MAX_INLINE_FILE_BYTES else ''}
        if ext=='.pdf' and len(data)<=MAX_INLINE_FILE_BYTES: payload['data_url']='data:application/pdf;base64,'+base64.b64encode(data).decode('ascii')
        return jsonify(payload)
    except Exception as e:return json_error(str(e),422,'file_parse_error')


def load_frames(data,ext):
    import pandas as pd
    if ext=='.csv':return {'CSV':pd.read_csv(io.BytesIO(data))}
    if ext=='.tsv':return {'TSV':pd.read_csv(io.BytesIO(data),sep='\t')}
    if ext in {'.xlsx','.xlsm','.xls'}:
        book=pd.ExcelFile(io.BytesIO(data));return {s:book.parse(s) for s in book.sheet_names}
    raise ValueError('Data Analysis supports CSV, TSV, XLSX, XLSM and XLS.')

def compute_profile(frames):
    import pandas as pd
    sheets={};total=0
    for sheet,df in frames.items():
        total+=len(df);num=df.select_dtypes(include='number');numeric={}
        if not num.empty:
            for col,row in num.describe().to_dict().items():numeric[str(col)]={k:(None if pd.isna(v) else float(v)) for k,v in row.items() if k in {'count','mean','std','min','25%','50%','75%','max'}}
        cats={}
        for col in df.select_dtypes(exclude='number').columns[:20]:
            vc=df[col].astype(str).replace('nan','').value_counts().head(5);cats[str(col)]={str(k):int(v) for k,v in vc.items()}
        sheets[sheet]={'rows':len(df),'columns':len(df.columns),'column_names':[str(c) for c in df.columns],'dtypes':{str(c):str(t) for c,t in df.dtypes.items()},'missing':{str(k):int(v) for k,v in df.isna().sum().items()},'duplicates':int(df.duplicated().sum()),'numeric_summary':numeric,'top_values':cats,'sample':df.head(12).fillna('').astype(str).to_dict(orient='records')}
    return {'total_rows':total,'sheets':sheets}

def compute_question(frames,q):
    import pandas as pd
    ql=q.lower();answers=[]
    for sheet,df in frames.items():
        cols={str(c).lower():c for c in df.columns}
        if 'duplicate' in ql: answers.append(f'{sheet}: {int(df.duplicated().sum())} duplicate rows.')
        if 'missing' in ql:
            miss=df.isna().sum().sort_values(ascending=False);top=miss[miss>0].head(5)
            answers.append(f"{sheet} missing: "+(', '.join(f'{c}={int(v)}' for c,v in top.items()) if len(top) else 'no missing values.'))
        for op,word in [('sum','sum'),('mean','average'),('mean','mean'),('max','maximum'),('min','minimum')]:
            if word in ql:
                candidates=[orig for low,orig in cols.items() if low in ql and pd.api.types.is_numeric_dtype(df[orig])]
                for col in candidates[:3]:
                    val=getattr(df[col],op)();answers.append(f'{sheet}: {op} of {col} = {float(val):.4f}.')
        if 'count' in ql and not any(x in ql for x in ['discount','account']): answers.append(f'{sheet}: {len(df)} rows.')
    return answers[:12]

@app.post('/api/data/analyze')
def data_analyze():
    f=request.files.get('file')
    if not f:return json_error('No CSV/Excel file provided')
    data=f.read();name=Path(f.filename or 'data').name;ext=Path(name).suffix.lower();question=request.form.get('question','').strip()
    if len(data)>20*1024*1024:return json_error('File exceeds the 20 MB upload limit',413,'file_too_large')
    try:
        frames=load_frames(data,ext);profile=compute_profile(frames);computed=compute_question(frames,question);answer=f'{name} contains {len(frames)} sheet(s) and {profile["total_rows"]} total rows.'
        if computed:answer+='\n\nComputed results:\n- '+'\n- '.join(computed)
        if question and GROQ_API_KEY:
            prompt=f'Use only the computed dataset facts below. Do not invent values. If the facts do not answer the question, say so.\nQUESTION: {question}\nCOMPUTED RESULTS: {json.dumps(computed,ensure_ascii=False)}\nPROFILE: {json.dumps(profile,ensure_ascii=False)[:50000]}'
            r=groq_complete([{'role':'system','content':'You are Vaani Data Analysis. Explain only grounded computed facts.'},{'role':'user','content':prompt}],False);answer=(r.choices[0].message.content or answer).strip()
        return jsonify({'ok':True,'name':name,'rows':profile['total_rows'],'sheets':len(frames),'columns':sum(len(x.columns) for x in frames.values()),'profile':profile,'computed':computed,'answer':answer})
    except Exception as e:return json_error(str(e),422,'data_analysis_error')

@app.get('/api/research')
def api_research():
    q=request.args.get('q','').strip()
    if not q:return json_error('Missing query')
    try:
        results=[]; seen=set()
        for qq in (q,f'{q} official source',f'{q} documentation'):
            try: found,_=search_web(qq,7)
            except Exception: continue
            for item in found:
                if item['url'] not in seen:
                    seen.add(item['url']); results.append(item)
        if not results:return json_error('Research search returned no usable sources.',502,'research_no_results')
        summary=''
        if GROQ_API_KEY:
            src='\n'.join(f"{x['title']}\n{x['text']}\n{x['url']}" for x in results[:15])
            r=groq_complete([{'role':'system','content':'Summarize only claims supported by the supplied search snippets. Distinguish uncertainty. Keep source URLs visible and do not invent facts.'},{'role':'user','content':f'Question: {q}\nSources:\n{src}'}],False)
            summary=(r.choices[0].message.content or '').strip()
        return jsonify({'ok':True,'query':q,'results':results[:15],'summary':summary,'source':'DuckDuckGo'})
    except Exception as e:return json_error(str(e),502,'research_error')


@app.post('/api/generate/image')
def generate_image():
    if not gemini_client:return json_error('GEMINI_API_KEY is not configured',503,'missing_gemini_key')
    data=request.get_json(silent=True) or {};prompt=str(data.get('prompt','')).strip();images=data.get('images') or []
    if not prompt:return json_error('Missing image prompt')
    try:
        contents=[prompt]
        for item in images[:3]:
            if item.get('data_url'):
                raw,mime=data_url_bytes(item['data_url']);contents.append(genai_types.Part.from_bytes(data=raw,mime_type=mime))
        response=gemini_client.models.generate_content(model=GEMINI_IMAGE_MODEL,contents=contents,config=genai_types.GenerateContentConfig(response_modalities=['IMAGE','TEXT']))
        for part in response.parts or []:
            inline=getattr(part,'inline_data',None)
            if inline and getattr(inline,'data',None):
                mime=getattr(inline,'mime_type',None) or 'image/png';return jsonify({'ok':True,'data_url':f'data:{mime};base64,{base64.b64encode(inline.data).decode("ascii")}','kind':'image','model':GEMINI_IMAGE_MODEL})
        return json_error('Gemini image model returned no image.',502,'empty_image_result')
    except Exception as e:
        app.logger.exception('image_generation_error')
        return json_error('Image creation is temporarily unavailable. Please try again in a moment.',503,'image_generation_error')

@app.post('/api/generate/edit')
def generate_edit():
    if not gemini_client:return json_error('GEMINI_API_KEY is not configured',503,'missing_gemini_key')
    data=request.get_json(silent=True) or {};prompt=str(data.get('prompt','')).strip();images=data.get('images') or []
    if not prompt:return json_error('Missing edit prompt')
    if not images or not images[0].get('data_url'):return json_error('Upload an image before editing')
    try:
        raw,mime=data_url_bytes(images[0]['data_url']);response=gemini_client.models.generate_content(model=GEMINI_IMAGE_MODEL,contents=[prompt,genai_types.Part.from_bytes(data=raw,mime_type=mime)],config=genai_types.GenerateContentConfig(response_modalities=['IMAGE','TEXT']))
        for part in response.parts or []:
            inline=getattr(part,'inline_data',None)
            if inline and getattr(inline,'data',None):
                outmime=getattr(inline,'mime_type',None) or 'image/png';return jsonify({'ok':True,'data_url':f'data:{outmime};base64,{base64.b64encode(inline.data).decode("ascii")}','kind':'image','model':GEMINI_IMAGE_MODEL})
        return json_error('Gemini image model returned no edited image.',502,'empty_edit_result')
    except Exception as e:
        app.logger.exception('image_edit_error')
        return json_error('Image editing is temporarily unavailable. Please try again in a moment.',503,'image_edit_error')

@app.post('/api/generate/video')
def generate_video():
    if not gemini_client:return json_error('GEMINI_API_KEY is not configured',503,'missing_gemini_key')
    data=request.get_json(silent=True) or {};prompt=str(data.get('prompt','')).strip()
    if not prompt:return json_error('Missing video prompt')
    try:
        operation=gemini_client.models.generate_videos(model=GEMINI_VIDEO_MODEL,prompt=prompt)
        import time
        while not operation.done:
            time.sleep(8);operation=gemini_client.operations.get(operation)
        video=operation.response.generated_videos[0]
        return jsonify({'ok':True,'video_file':getattr(getattr(video,'video',None),'uri',None),'kind':'video','model':GEMINI_VIDEO_MODEL,'note':'Video generation completed. Use the returned Gemini file URI with the Gemini Files API.'})
    except Exception as e:
        app.logger.exception('video_generation_error')
        return json_error('Video creation is temporarily unavailable. Please try again in a moment.',503,'video_generation_error')

def route_model(data):
    return GROQ_MODEL

@app.post('/api/chat')
def chat():
    data=request.get_json(silent=True) or {};message=str(data.get('message','')).strip()
    if not message and not has_attachments(data.get('live_context') or {}):return json_error('Message or attachment is required')
    live=data.get('live_context') or {};history=data.get('history') or [];mode=data.get('mode','General')
    try:
        if has_attachments(live):
            image_count=len(live.get('images') or [])
            attachment_count=len(live.get('attachments') or [])
            app.logger.info('Vision request: images=%s attachments=%s',image_count,attachment_count)
            text,vision_model=gemini_complete(message or 'Inspect the attached file/image and explain what is important.',history,mode,live)
            return jsonify({'ok':True,'text':text,'model_used':vision_model,'vision_routed':True,'vision_image_count':len(live.get('images') or []),'vision_attachment_count':len(live.get('attachments') or [])})
        r=groq_complete(build_text_messages(message,history,mode,live),False)
        return jsonify({'ok':True,'text':r.choices[0].message.content or '','provider':'groq','model_used':GROQ_MODEL,'vision_routed':False})
    except Exception as e:
        app.logger.exception('Chat request failed')
        msg=str(e)
        safe=msg if msg.startswith(('Image analysis is','GROQ_API_KEY is')) else 'Vaani could not complete that request right now. Please try again.'
        return json_error(safe,503 if 'Image analysis is' in safe else 502,'chat_error')

@app.post('/api/chat/stream')
def chat_stream():
    data=request.get_json(silent=True) or {};message=str(data.get('message','')).strip();live=data.get('live_context') or {};history=data.get('history') or [];mode=data.get('mode','General')
    if has_attachments(live):return jsonify({'ok':False,'error':'Multimodal requests use the non-streaming Gemini route.','code':'multimodal_use_chat'}),409
    if not message:return json_error('Message is required')
    def generate():
        try:
            stream=groq_complete(build_text_messages(message,history,mode,live),True)
            for chunk in stream:
                delta=chunk.choices[0].delta.content if chunk.choices else None
                if delta:yield 'data: '+json.dumps({'text':delta})+'\n\n'
            yield 'data: '+json.dumps({'meta':{'model_used':GROQ_MODEL,'provider':'groq','vision_routed':False}})+'\n\n';yield 'data: [DONE]\n\n'
        except Exception as e:
            yield 'data: '+json.dumps({'error':str(e),'code':'chat_stream_error'})+'\n\n';yield 'data: [DONE]\n\n'
    return Response(stream_with_context(generate()),mimetype='text/event-stream',headers={'Cache-Control':'no-cache, no-transform','X-Accel-Buffering':'no','Connection':'keep-alive'})

@app.get('/api/models')
def models():
    return jsonify({'ok':True,'selected_supported':[GROQ_MODEL],'models':[{'id':GROQ_MODEL,'modalities':['text']}],'vision':{'available':bool(gemini_client),'modalities':['text','image','video','audio','pdf'],'fallback_count':len(GEMINI_VISION_FALLBACKS)}})

@app.errorhandler(413)
def too_large(_):return json_error('Request/file exceeds the 20 MB limit.',413,'file_too_large')
@app.errorhandler(500)
def internal(_):return json_error('Unexpected backend error.',500,'internal_error')

if __name__=='__main__':
    print('Vaani backend '+APP_VERSION+':',BASE_DIR);print('Groq configured:',bool(GROQ_API_KEY),'Gemini configured:',bool(GEMINI_API_KEY));app.run(host='127.0.0.1',port=int(os.getenv('PORT','5288')),debug=True)
